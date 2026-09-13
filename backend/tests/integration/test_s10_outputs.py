"""Stage 10 산출물 — 풋볼필드·감사추적·케이스 JSON·워크북·리포트."""

from __future__ import annotations

import io
import json

from openpyxl import load_workbook

from app.api import service
from app.core import stage_gate
from app.db.models import (
    AssetResult,
    Case,
    DcfResult,
    FinalValue,
    FinancialsNormalized,
    Lineage,
    MarketResult,
    McResult,
    SourceType,
)
from app.excel import exporters
from app.reports import report


def _seed(db) -> Case:
    case = Case(case_id="c10", corp_code="00126380", corp_name="삼성전자", stock_code="005930",
                induty_code="264", valuation_date="2023-12-28", purpose="general_ma")
    db.add(case)
    db.add(DcfResult(case_id="c10", payload={"equity_value": 1000.0, "tv_ratio": 0.6,
                                             "rows": [{"t": 1, "revenue": 100, "fcff": 10, "pv": 9}]}))
    db.add(McResult(case_id="c10", payload={"mean": 1000.0, "p10": 800.0, "p50": 1000.0,
                                            "p90": 1250.0}))
    db.add(MarketResult(case_id="c10", payload={"range": {"min": 1100.0, "median": 1200.0,
                                                          "max": 1300.0}, "peers": []}))
    db.add(AssetResult(case_id="c10", payload={"adjusted_net_asset": 900.0,
                                               "value_with_premium": 900.0}))
    db.add(FinalValue(case_id="c10", payload={
        "synthesis": {"mode": "weighted", "final": 1033.0},
        "flags": [{"key": "peers_min", "triggered": True, "severity": "warn",
                   "description": "유사기업 < 5개", "value": 1, "threshold": 5}],
        "flag_approvals": {}}))
    db.add(FinancialsNormalized(case_id="c10", year=2023, account="revenue", value=100.0))
    db.add(Lineage(case_id="c10", field_name="wacc.rf", value=3.18, source_type=SourceType.API,
                   source_ref="ECOS"))
    db.add(Lineage(case_id="c10", field_name="fin.2023.dna", value=40000.0,
                   source_type=SourceType.EXPERT_OVERRIDE, system_value=None,
                   rationale="주석 감가상각비", changed_by="analyst"))
    for s in range(10):
        stage_gate.mark_reviewed(db, "c10", s)
        stage_gate.approve(db, "c10", s, approved_by="a")
    db.flush()
    return case


def test_football_field_uses_mc_range(db):
    case = _seed(db)
    ff = {f["method"]: f for f in service.football_field(db, case)}
    # 수익가치는 MC P10~P90
    assert ff["수익가치(DCF)"]["low"] == 800.0 and ff["수익가치(DCF)"]["high"] == 1250.0
    assert ff["시장가치"]["low"] == 1100.0 and ff["시장가치"]["high"] == 1300.0
    assert ff["자산가치"]["base"] == 900.0


def test_summary_and_final(db):
    case = _seed(db)
    s = service.stage10_summary(db, case)
    assert s["final_value"] == 1033.0
    assert len(s["football_field"]) == 3
    assert s["stage_status"][9] == "APPROVED"


def test_lineage_rows_include_override(db):
    case = _seed(db)
    rows = service.lineage_rows(db, case)
    ov = [r for r in rows if r["source_type"] == "EXPERT_OVERRIDE"]
    assert ov and ov[0]["rationale"] == "주석 감가상각비"


def test_case_json_snapshot(db):
    case = _seed(db)
    cj = service.export_case_json(db, case)
    assert cj["case"]["corp_name"] == "삼성전자"
    assert cj["dcf_result"]["equity_value"] == 1000.0
    assert cj["final_value"]["synthesis"]["final"] == 1033.0
    # JSON 직렬화 가능
    json.dumps(cj, ensure_ascii=False)


def test_full_workbook_sheets(db):
    case = _seed(db)
    cj = service.export_case_json(db, case)
    summ = service.stage10_summary(db, case)
    wb = load_workbook(io.BytesIO(exporters.full_workbook(cj, summ)))
    for sheet in ("요약", "정규화재무", "WACC", "DCF", "최종", "감사추적"):
        assert sheet in wb.sheetnames


def test_html_report_renders_korean(db):
    case = _seed(db)
    cj = service.export_case_json(db, case)
    summ = service.stage10_summary(db, case)
    html = report.render_report(cj, summ)
    assert "삼성전자" in html
    assert "풋볼필드" in html
    assert "1,033" in html  # 최종가치 포맷
    assert "주석 감가상각비" in html  # 오버라이드 내역
