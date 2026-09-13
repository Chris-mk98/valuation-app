"""excel/ 라운드트립·검증 테스트."""

from __future__ import annotations

import io

from openpyxl import load_workbook

from app.excel.exporters import stage1_raw_dump, stage2_normalized
from app.excel.importers import parse_account_map
from app.excel.templates import account_map_template
from app.params.account_map import load_account_map
from app.stages.s2_normalize import normalize_year


def test_account_map_template_roundtrip():
    original = load_account_map()
    xlsx = account_map_template(original)
    mapping, errors = parse_account_map(xlsx)
    assert errors == []
    # 주요 섹션·항목 보존
    assert "revenue" in mapping["targets"]
    assert mapping["targets"]["revenue"]["account_ids"] == ["ifrs-full_Revenue"]
    assert "purchase_ppe" in mapping["components"]
    assert "dna_combined" in mapping["dna_candidates"]
    # 재구성한 매핑으로 정규화가 동작(라운드트립 보장)
    rows = [
        {"account_id": "ifrs-full_Revenue", "account_nm": "매출액", "sj_div": "IS",
         "thstrm_amount": "300000000"},
    ]
    res = normalize_year(rows, mapping)
    assert res["accounts"]["revenue"] == 300.0  # 300,000,000원 ÷ 1e6 = 300 백만원


def test_importer_reports_cell_errors():
    # 잘못된 section 과 빈 internal 을 포함한 워크북 생성
    from openpyxl import Workbook
    wb = Workbook()
    ws = wb.active
    ws.title = "account_map"
    ws.append(["section", "internal", "statement", "account_ids", "account_names"])
    ws.append(["targets", "revenue", "IS", "ifrs-full_Revenue", "매출액"])  # 정상
    ws.append(["WRONG", "x", "IS", "a", "b"])                              # section 오류
    ws.append(["targets", "", "IS", "a", "b"])                             # internal 누락
    ws.append(["targets", "ebit", "XX", "a", "b"])                         # statement 오류
    buf = io.BytesIO()
    wb.save(buf)
    mapping, errors = parse_account_map(buf.getvalue())
    cols = {e["column"] for e in errors}
    assert "section" in cols and "internal" in cols and "statement" in cols
    assert "revenue" in mapping["targets"]  # 정상 행은 반영


def test_stage2_export_readable():
    normalized = {
        "years": [2022, 2023],
        "accounts": {2022: {"revenue": 100.0, "ebit": 10.0}, 2023: {"revenue": 120.0, "ebit": 12.0}},
        "unmapped": {2022: ["dna"], 2023: ["dna"]},
        "warnings": {2022: [], 2023: []},
    }
    xlsx = stage2_normalized(normalized)
    wb = load_workbook(io.BytesIO(xlsx))
    ws = wb["정규화재무(백만원)"]
    assert ws.cell(row=1, column=1).value == "account"
    assert ws.cell(row=1, column=3).value == "2023"
    # revenue 는 필수계정 첫 행
    assert ws.cell(row=2, column=1).value == "revenue"
    assert ws.cell(row=2, column=3).value == 120.0


def test_stage1_raw_dump_sheets():
    raw = {2023: [{"sj_div": "IS", "account_id": "ifrs-full_Revenue", "account_nm": "매출액",
                   "thstrm_amount": "300000000"}]}
    xlsx = stage1_raw_dump(raw)
    wb = load_workbook(io.BytesIO(xlsx))
    assert "2023" in wb.sheetnames
    assert wb["2023"].cell(row=2, column=3).value == "매출액"
