"""업로드 템플릿 생성 (CLAUDE.md 원칙 3: 입력이 필요한 표는 템플릿 다운로드 제공).

account_map_template: 계정 매핑표를 평면 테이블로 내보낸다(라운드트립용).
컬럼: section | internal | statement | account_ids | account_names
"""

from __future__ import annotations

import io

from openpyxl import Workbook
from openpyxl.styles import Font

MAP_COLUMNS = ["section", "internal", "statement", "account_ids", "account_names"]
_SECTIONS = ["targets", "components", "dna_candidates"]


def account_map_template(mapping: dict) -> bytes:
    """현재 매핑을 채운 템플릿 워크북(bytes). 사용자가 수정 후 업로드."""
    wb = Workbook()
    ws = wb.active
    ws.title = "account_map"
    ws.append(MAP_COLUMNS)
    for c in ws[1]:
        c.font = Font(bold=True)

    for section in _SECTIONS:
        for internal, spec in (mapping.get(section) or {}).items():
            ws.append([
                section,
                internal,
                spec.get("statement", ""),
                ", ".join(spec.get("account_ids") or []),
                ", ".join(spec.get("account_names") or []),
            ])

    # 안내 시트
    guide = wb.create_sheet("설명")
    guide.append(["section 은 targets/components/dna_candidates 중 하나."])
    guide.append(["account_ids, account_names 는 쉼표로 구분. account_id 우선 매칭, 없으면 account_nm."])
    guide.append(["statement 는 BS/IS/CF/CIS. 수정 후 이 파일을 그대로 업로드하세요."])

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
