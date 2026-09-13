"""업로드 검증·반영 (CLAUDE.md 원칙 3: 검증 후 셀 단위 오류 표시).

parse_account_map: account_map 템플릿(xlsx)을 매핑 dict 로 파싱하고 셀 단위 오류를 수집.
"""

from __future__ import annotations

import io

from openpyxl import load_workbook

from app.excel.templates import MAP_COLUMNS

_VALID_SECTIONS = {"targets", "components", "dna_candidates"}
_VALID_STATEMENTS = {"BS", "IS", "CF", "CIS"}


class ImportError_(ValueError):
    """검증 실패(치명적)."""


def _split(cell: object) -> list[str]:
    if cell is None:
        return []
    return [x.strip() for x in str(cell).split(",") if x.strip()]


def parse_account_map(file_bytes: bytes) -> tuple[dict, list[dict]]:
    """xlsx 바이트 → (mapping dict, errors).

    errors: [{"row": int, "column": str, "message": str}]  (비어 있으면 검증 통과)
    mapping: {"unit_divisor": ..., "targets": {...}, "components": {...}, "dna_candidates": {...}}
    """
    wb = load_workbook(io.BytesIO(file_bytes), read_only=True, data_only=True)
    if "account_map" not in wb.sheetnames:
        raise ImportError_("시트 'account_map' 가 없습니다.")
    ws = wb["account_map"]

    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        raise ImportError_("빈 시트입니다.")

    header = [str(c).strip() if c is not None else "" for c in rows[0]]
    if header[: len(MAP_COLUMNS)] != MAP_COLUMNS:
        raise ImportError_(f"헤더가 템플릿과 다릅니다. 기대: {MAP_COLUMNS}, 실제: {header}")

    mapping: dict = {"unit_divisor": 1_000_000, "targets": {}, "components": {},
                     "dna_candidates": {}}
    errors: list[dict] = []
    seen: set[tuple[str, str]] = set()

    for i, row in enumerate(rows[1:], start=2):  # 엑셀 행번호(1=헤더)
        section, internal, statement, ids, names = (list(row) + [None] * 5)[:5]
        if section is None and internal is None:
            continue  # 빈 행 스킵
        section = str(section).strip() if section else ""
        internal = str(internal).strip() if internal else ""
        statement = str(statement).strip() if statement else ""

        if section not in _VALID_SECTIONS:
            errors.append({"row": i, "column": "section",
                           "message": f"section 은 {sorted(_VALID_SECTIONS)} 중 하나여야 합니다."})
            continue
        if not internal:
            errors.append({"row": i, "column": "internal", "message": "internal 은 필수입니다."})
            continue
        if statement not in _VALID_STATEMENTS:
            errors.append({"row": i, "column": "statement",
                           "message": f"statement 는 {sorted(_VALID_STATEMENTS)} 중 하나."})
            continue
        if (section, internal) in seen:
            errors.append({"row": i, "column": "internal",
                           "message": f"중복 항목: {section}/{internal}"})
            continue
        seen.add((section, internal))

        account_ids = _split(ids)
        account_names = _split(names)
        if not account_ids and not account_names:
            errors.append({"row": i, "column": "account_ids",
                           "message": "account_ids 또는 account_names 중 하나는 필요합니다."})
            continue

        mapping[section][internal] = {
            "statement": statement,
            "account_ids": account_ids,
            "account_names": account_names,
        }

    return mapping, errors


def parse_adjustments(file_bytes: bytes) -> tuple[list[dict], list[dict]]:
    """자산조정 템플릿(xlsx) → (adjustments, errors). 컬럼 label, amount, rationale."""
    wb = load_workbook(io.BytesIO(file_bytes), read_only=True, data_only=True)
    if "adjustments" not in wb.sheetnames:
        raise ImportError_("시트 'adjustments' 가 없습니다.")
    rows = list(wb["adjustments"].iter_rows(values_only=True))
    if not rows:
        raise ImportError_("빈 시트입니다.")
    header = [str(c).strip() if c is not None else "" for c in rows[0]]
    if header[:3] != ["label", "amount", "rationale"]:
        raise ImportError_("헤더가 템플릿과 다릅니다. 기대: ['label','amount','rationale']")

    adjustments: list[dict] = []
    errors: list[dict] = []
    for i, row in enumerate(rows[1:], start=2):
        label, amount, rationale = (list(row) + [None] * 3)[:3]
        if label is None and amount is None:
            continue
        if not label:
            errors.append({"row": i, "column": "label", "message": "label 필수"})
            continue
        try:
            amt = float(amount)
        except (TypeError, ValueError):
            errors.append({"row": i, "column": "amount", "message": "숫자가 아닙니다."})
            continue
        if not rationale or not str(rationale).strip():
            errors.append({"row": i, "column": "rationale", "message": "사유 필수"})
            continue
        adjustments.append({"label": str(label), "amount": amt, "rationale": str(rationale)})
    return adjustments, errors


def parse_transactions(file_bytes: bytes) -> tuple[list[dict], list[dict]]:
    """거래사례 템플릿(xlsx) → (transactions, errors). multiple 미기재 시 deal_ev/ebitda 산출."""
    wb = load_workbook(io.BytesIO(file_bytes), read_only=True, data_only=True)
    if "transactions" not in wb.sheetnames:
        raise ImportError_("시트 'transactions' 가 없습니다.")
    rows = list(wb["transactions"].iter_rows(values_only=True))
    if not rows:
        raise ImportError_("빈 시트입니다.")

    header = [str(c).strip() if c is not None else "" for c in rows[0]]
    expected = ["target_name", "deal_ev", "ebitda", "multiple"]
    if header[:4] != expected:
        raise ImportError_(f"헤더가 템플릿과 다릅니다. 기대: {expected}")

    txns: list[dict] = []
    errors: list[dict] = []
    for i, row in enumerate(rows[1:], start=2):
        name, deal_ev, ebitda, multiple = (list(row) + [None] * 4)[:4]
        if name is None and deal_ev is None and ebitda is None:
            continue
        try:
            ev = float(deal_ev) if deal_ev is not None else None
            eb = float(ebitda) if ebitda is not None else None
            mult = float(multiple) if multiple not in (None, "") else None
        except (TypeError, ValueError):
            errors.append({"row": i, "column": "deal_ev/ebitda/multiple",
                           "message": "숫자가 아닙니다."})
            continue
        if mult is None:
            if ev is None or not eb:
                errors.append({"row": i, "column": "multiple",
                               "message": "multiple 또는 (deal_ev & ebitda)가 필요합니다."})
                continue
            mult = ev / eb
        txns.append({"target_name": str(name) if name else "", "deal_ev": ev,
                     "ebitda": eb, "multiple": mult})
    return txns, errors
